"""Accept adaptable movie slides through source inspection and ODP conversion."""

# Standard Library
import sys
import pathlib
import tempfile

# PIP3 modules
import pptx
import PIL.Image #pillow


TESTS_DIR = pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0, str(TESTS_DIR))

# local repo modules
import file_utils

REPO_ROOT = pathlib.Path(file_utils.get_repo_root())

# local repo modules
import slide_maker.moviedata
import slide_maker.slide_builder
import slide_maker.slide_convert


OUTPUT_DIRECTORY = REPO_ROOT / "output_smoke" / "adaptive_slides"
OUTPUT_PATH = OUTPUT_DIRECTORY / "adaptive_slides.odp"


#============================================
def require(condition: bool, message: str) -> None:
	"""Raise a clear E2E failure when adaptive output is not student-ready."""
	if condition:
		return
	raise RuntimeError(message)


#============================================
def shape_by_name(slide: object, name: str) -> object | None:
	"""Return one semantic role when it remains in an adaptive movie slide."""
	candidates = [shape for shape in slide.shapes if shape.name == name]
	require(len(candidates) <= 1, f"Built slide has an ambiguous semantic role: {name}")
	if not candidates:
		return None
	return candidates[0]


#============================================
def movie_cases(poster_path: pathlib.Path) -> tuple[tuple[str, slide_maker.moviedata.MovieData], ...]:
	"""Return inline records that cover available and absent movie metadata."""
	complete = slide_maker.moviedata.MovieData(
		title="Complete Metadata",
		year=2026,
		plot="A complete record establishes the reference slide.",
		genres=["Drama"],
		runtime_minutes=101,
		directors=["Ada Director"],
		imdb_rating=8.1,
		imdb_votes=1200,
		rt_tomatometer=84,
		rt_audience_score=90,
		rt_state="fresh",
		rt_consensus="A concise review summary.",
		metascore=55,
		metascore_band="middle",
		poster_path=str(poster_path),
		original_language="English",
	)
	non_english = slide_maker.moviedata.MovieData(
		title="Subtitled Feature",
		year=2024,
		plot="A language-bearing movie needs no ratings to be useful.",
		runtime_minutes=94,
		poster_path=str(poster_path),
		original_language="Korean",
	)
	poster_free = slide_maker.moviedata.MovieData(
		title="Poster-Free Feature",
		year=2023,
		genres=["Documentary"],
		directors=["Bea Director"],
		runtime_minutes=70,
		original_language="English",
	)
	title_only = slide_maker.moviedata.MovieData(title="Title Only")
	zero_scores = slide_maker.moviedata.MovieData(
		title="Zero Scores",
		year=2022,
		imdb_rating=0.0,
		imdb_votes=0,
		rt_tomatometer=0,
		rt_audience_score=0,
		rt_state="rotten",
		metascore=0,
		metascore_band="low",
		poster_path=str(poster_path),
	)
	missing_votes = slide_maker.moviedata.MovieData(
		title="Votes Unavailable",
		year=2021,
		imdb_rating=7.2,
		poster_path=str(poster_path),
	)
	long_text = slide_maker.moviedata.MovieData(
		title="Long Text",
		year=2020,
		plot=(
			"A carefully written plot gives students enough context to decide whether a film "
			"belongs in their study plan while remaining readable on a lecture slide."
		),
		rt_consensus=(
			"The review summary also uses enough words to exercise text fitting without requiring "
			"any missing scores, identifiers, or artificial placeholder labels."
		),
		poster_path=str(poster_path),
		original_language="Spanish",
	)
	return (
		("complete", complete),
		("non_english", non_english),
		("poster_free", poster_free),
		("title_only", title_only),
		("zero_scores", zero_scores),
		("missing_votes", missing_votes),
		("long_text", long_text),
	)


#============================================
def expected_fragments(movie_data: slide_maker.moviedata.MovieData) -> tuple[str, ...]:
	"""Return text that is present only when its source data is present."""
	fragments = []
	if movie_data.plot:
		fragments.append(movie_data.plot)
	if movie_data.imdb_rating is not None:
		fragments.append(f"IMDB rating {movie_data.imdb_rating:.1f}")
	if movie_data.imdb_votes is not None:
		votes = slide_maker.slide_builder.format_compact_count(movie_data.imdb_votes)
		fragments.append(f"{votes} votes")
	if movie_data.rt_tomatometer is not None:
		fragments.append(f"{movie_data.rt_tomatometer}%")
	if movie_data.rt_audience_score is not None:
		fragments.append("Audience:")
	if movie_data.metascore is not None:
		fragments.append("MS")
	if movie_data.genres:
		fragments.append(f"Genre: {', '.join(movie_data.genres)}")
	if movie_data.directors:
		fragments.append(f"Director: {', '.join(movie_data.directors)}")
	if movie_data.runtime_minutes is not None:
		fragments.append(f"Run time: {movie_data.runtime_minutes} min")
	if movie_data.original_language:
		fragments.append(f"Original language: {movie_data.original_language}")
	if movie_data.rt_consensus:
		fragments.append(f"Review Summary: {movie_data.rt_consensus}")
	return tuple(fragments)


#============================================
def validate_source_slide(
	slide: object,
	template_slide: object,
	movie_data: slide_maker.moviedata.MovieData,
) -> None:
	"""Require source roles and text to reflect exactly available metadata."""
	title = shape_by_name(slide, slide_maker.slide_builder.TEMPLATE_TITLE_NAME)
	require(title is not None, "Adaptive slide is missing its title role")
	require(
		title.text == slide_maker.slide_builder.movie_title(movie_data),
		"Adaptive title did not preserve the available title and year",
	)
	outline = shape_by_name(slide, slide_maker.slide_builder.TEMPLATE_OUTLINE_NAME)
	poster = shape_by_name(slide, slide_maker.slide_builder.TEMPLATE_POSTER_NAME)
	fragments = expected_fragments(movie_data)
	require((outline is not None) == bool(fragments), "Adaptive outline presence did not match source text")
	require((poster is not None) == (movie_data.poster_path is not None), "Adaptive poster presence changed")
	if outline is not None:
		for fragment in fragments:
			require(fragment in outline.text, f"Adaptive outline omitted {fragment!r}")
		require("N/A" not in outline.text, "Adaptive outline invented an N/A placeholder")
		if movie_data.poster_path is None:
			template_outline = shape_by_name(template_slide, slide_maker.slide_builder.TEMPLATE_OUTLINE_NAME)
			require(template_outline is not None, "Template has no outline anchor")
			require(outline.width > template_outline.width, "Poster-free outline did not use poster space")
	if movie_data.imdb_rating is not None and movie_data.imdb_votes is None:
		require("votes" not in outline.text, "Missing vote count produced a placeholder")
	if movie_data.title == "Subtitled Feature":
		for label in ("IMDB rating", "Critics:", "Audience:", "MS "):
			require(label not in outline.text, "No-rating slide invented a ratings label")
	for shape in (title, outline):
		if shape is None:
			continue
		for paragraph in shape.text_frame.paragraphs:
			for run in paragraph.runs:
				require(
					run.font.name == slide_maker.slide_builder.FONT_NAME,
					"Adaptive source text does not use Atkinson Hyperlegible Next",
				)


#============================================
def main() -> None:
	"""Build and convert a persisted deck covering sparse movie-slide states."""
	template_path = REPO_ROOT / "template" / "movie_slide_template.pptx"
	require(template_path.is_file(), f"Movie slide template is absent: {template_path}")
	OUTPUT_DIRECTORY.mkdir(parents=True, exist_ok=True)
	with tempfile.TemporaryDirectory() as temporary_directory:
		work_dir = pathlib.Path(temporary_directory)
		poster_path = work_dir / "runtime_poster.png"
		PIL.Image.new("RGB", (600, 900), color=(35, 72, 115)).save(poster_path)
		cases = movie_cases(poster_path)
		scratch_path = work_dir / "adaptive_slides.pptx"
		presentation = pptx.Presentation(template_path)
		for unused_name, movie_data in cases:
			slide_maker.slide_builder.append_movie_slide(presentation, movie_data)
		presentation.save(scratch_path)
		require(scratch_path.is_file(), "Adaptive source deck was not written")
		source = pptx.Presentation(scratch_path)
		template_slide = source.slides[0]
		visible_slides = [slide for slide in source.slides if slide._element.get("show") != "0"]
		require(len(visible_slides) == len(cases), "Adaptive source deck changed the case count")
		for slide, (case_name, movie_data) in zip(visible_slides, cases, strict=True):
			validate_source_slide(slide, template_slide, movie_data)
			print(f"Adaptive source acceptance passed: {case_name}")
		slide_maker.slide_convert.convert_presentation(scratch_path, OUTPUT_PATH, len(cases))
	require(OUTPUT_PATH.is_file(), f"Adaptive ODP is absent: {OUTPUT_PATH}")
	require(OUTPUT_PATH.stat().st_size > 0, "Adaptive ODP is empty")
	print(f"Adaptive-slide E2E passed: {OUTPUT_PATH}")


if __name__ == "__main__":
	main()
